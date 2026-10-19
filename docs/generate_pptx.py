# -*- coding: utf-8 -*-
"""Generate SFTP PoC overview PowerPoint."""

from pptx import Presentation
from pptx.util import Inches, Pt
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_SHAPE

OUTPUT = r"D:\Cursor\SFTP\SFTP_PoC_개요.pptx"

# Light (white) background + high-contrast text for readability
BG = RGBColor(0xFF, 0xFF, 0xFF)
ACCENT = RGBColor(0x0B, 0x5F, 0xAD)      # strong blue
TEXT = RGBColor(0x1A, 0x1A, 0x1A)        # near-black body text
MUTED = RGBColor(0x44, 0x44, 0x44)       # dark gray secondary
CARD = RGBColor(0xF3, 0xF6, 0xFA)        # light card fill
CARD_LINE = RGBColor(0xC5, 0xD0, 0xDC)
TITLE = RGBColor(0x0A, 0x2E, 0x5C)       # navy titles
WARN = RGBColor(0xB5, 0x5A, 0x00)        # dark amber


def set_run(run, size=18, bold=False, color=TEXT):
    run.font.size = Pt(size)
    run.font.bold = bold
    run.font.color.rgb = color
    run.font.name = "Malgun Gothic"


def add_bg(slide):
    shape = slide.shapes.add_shape(
        MSO_SHAPE.RECTANGLE, Inches(0), Inches(0), Inches(13.333), Inches(7.5)
    )
    shape.fill.solid()
    shape.fill.fore_color.rgb = BG
    shape.line.fill.background()


def add_accent_bar(slide):
    bar = slide.shapes.add_shape(
        MSO_SHAPE.RECTANGLE, Inches(0), Inches(0), Inches(0.18), Inches(7.5)
    )
    bar.fill.solid()
    bar.fill.fore_color.rgb = ACCENT
    bar.line.fill.background()


def add_title(slide, text, top=0.35, size=32):
    box = slide.shapes.add_textbox(Inches(0.55), Inches(top), Inches(12.2), Inches(0.8))
    tf = box.text_frame
    tf.clear()
    p = tf.paragraphs[0]
    run = p.add_run()
    run.text = text
    set_run(run, size=size, bold=True, color=TITLE)
    return box


def add_bullets(slide, items, left=0.55, top=1.3, width=12.2, height=5.5, size=18):
    box = slide.shapes.add_textbox(Inches(left), Inches(top), Inches(width), Inches(height))
    tf = box.text_frame
    tf.word_wrap = True
    tf.clear()
    for i, item in enumerate(items):
        p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        p.level = item[0] if isinstance(item, tuple) else 0
        text = item[1] if isinstance(item, tuple) else item
        run = p.add_run()
        run.text = text
        set_run(run, size=size, bold=False, color=TEXT)
        p.space_after = Pt(8)
    return box


def add_footer(slide, page, total):
    box = slide.shapes.add_textbox(Inches(0.55), Inches(7.05), Inches(12.2), Inches(0.3))
    tf = box.text_frame
    tf.clear()
    p = tf.paragraphs[0]
    run = p.add_run()
    run.text = f"Windows SFTP Download Client PoC  |  {page}/{total}"
    set_run(run, size=11, color=MUTED)


def add_card(slide, x, y, w, h, title, body_lines):
    shape = slide.shapes.add_shape(
        MSO_SHAPE.ROUNDED_RECTANGLE, Inches(x), Inches(y), Inches(w), Inches(h)
    )
    shape.fill.solid()
    shape.fill.fore_color.rgb = CARD
    shape.line.color.rgb = CARD_LINE
    shape.line.width = Pt(1.25)
    try:
        shape.adjustments[0] = 0.08
    except Exception:
        pass

    t = slide.shapes.add_textbox(Inches(x + 0.2), Inches(y + 0.15), Inches(w - 0.4), Inches(0.4))
    tf = t.text_frame
    tf.clear()
    p = tf.paragraphs[0]
    run = p.add_run()
    run.text = title
    set_run(run, size=16, bold=True, color=ACCENT)

    b = slide.shapes.add_textbox(Inches(x + 0.2), Inches(y + 0.55), Inches(w - 0.4), Inches(h - 0.7))
    tf = b.text_frame
    tf.word_wrap = True
    tf.clear()
    for i, line in enumerate(body_lines):
        p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        run = p.add_run()
        run.text = line
        set_run(run, size=14, color=TEXT)
        p.space_after = Pt(4)


def main():
    prs = Presentation()
    prs.slide_width = Inches(13.333)
    prs.slide_height = Inches(7.5)
    blank = prs.slide_layouts[6]
    slides_meta = []

    def new_slide():
        s = prs.slides.add_slide(blank)
        add_bg(s)
        add_accent_bar(s)
        slides_meta.append(s)
        return s

    # 1 Title
    s = new_slide()
    box = s.shapes.add_textbox(Inches(0.7), Inches(2.2), Inches(11.5), Inches(1.2))
    tf = box.text_frame
    tf.clear()
    p = tf.paragraphs[0]
    run = p.add_run()
    run.text = "Windows SFTP Download Client PoC"
    set_run(run, size=40, bold=True, color=TITLE)

    box = s.shapes.add_textbox(Inches(0.7), Inches(3.5), Inches(11.5), Inches(1.2))
    tf = box.text_frame
    tf.clear()
    p = tf.paragraphs[0]
    run = p.add_run()
    run.text = "전체 요구사항 · 라이브러리 선정 · 개발 진행 · 서버 운용"
    set_run(run, size=22, color=MUTED)

    box = s.shapes.add_textbox(Inches(0.7), Inches(5.2), Inches(11.5), Inches(0.8))
    tf = box.text_frame
    tf.clear()
    p = tf.paragraphs[0]
    run = p.add_run()
    run.text = "2026-07-29  |  C++ / Visual Studio 2015  |  libssh2 + WinCNG"
    set_run(run, size=16, bold=True, color=ACCENT)

    # 2 Goal
    s = new_slide()
    add_title(s, "1. 프로젝트 목적")
    add_bullets(
        s,
        [
            "Linux SFTP 서버 → Windows 로컬로 사전 지정 파일 자동 다운로드",
            "사용자 조작 없음 (GUI / 목록 조회 / 업로드 / 파일 선택 없음)",
            "배포: Client.exe + config.ini (목표 DLL 없음, 10MB 미만)",
            "고객 PC Runtime·별도 SW 설치 금지 (/MT 정적 CRT)",
            "Host Key SHA256 Fingerprint 검증 필수 (미등록·불일치 시 연결 금지)",
        ],
    )

    # 3 Requirements
    s = new_slide()
    add_title(s, "2. 핵심 요구사항")
    add_card(s, 0.55, 1.3, 4.0, 5.2, "개발 환경", [
        "C++ Console Application",
        "VS2015 (v140) x64",
        "Windows 10/11",
        "C#/.NET/Python 제품 금지",
        "정적 링크 우선",
    ])
    add_card(s, 4.75, 1.3, 4.0, 5.2, "보안", [
        "SSH/SFTP 직접 구현 금지",
        "검증된 라이브러리 사용",
        "Host Key 검증 생략 금지",
        "trust-all 코드 금지",
        "Fingerprint 불일치 즉시 종료",
    ])
    add_card(s, 8.95, 1.3, 3.8, 5.2, "오류 표시", [
        "서버 연결 실패",
        "Host Key 불일치/미등록",
        "로그인 실패",
        "RemoteFile 없음",
        "LocalFile 생성 실패",
        "다운로드 실패/성공",
    ])

    # 4 Flow
    s = new_slide()
    add_title(s, "3. 자동 실행 흐름 (12단계)")
    add_bullets(
        s,
        [
            "1) config.ini 읽기  →  2) TCP/SSH 연결  →  3) Host Key 수신",
            "4) SHA-256 Fingerprint 계산  →  5) 설정값 비교 (불일치/미등록 시 종료)",
            "6) ID/PW 인증  →  7) SFTP 세션 시작",
            "8) RemoteFile 존재 확인  →  9) LocalFile 다운로드",
            "10) 성공 검증  →  11) 연결 정상 종료  →  12) 프로그램 종료",
            "파일 대상: RemoteFileN / LocalFileN (1~20개)",
        ],
        size=18,
    )

    # 5 Library
    s = new_slide()
    add_title(s, "4. 라이브러리 선정: libssh2 + WinCNG")
    add_bullets(
        s,
        [
            "선정: libssh2 1.11.1 (BSD-3-Clause) + Windows WinCNG",
            "이유: 정적 링크 가능, OpenSSL DLL 불필요, SHA256 Host Key API 제공",
            "대안 제외: libssh(LGPL), libcurl(의존성 증가), OpenSSH 바이너리 호출",
            "제한: Ed25519 host key 미지원 → 서버 RSA/ECDSA host key 필요",
            "라이선스 고지: dist/LICENSE_libssh2.txt 동봉",
        ],
    )

    # 6 Progress
    s = new_slide()
    add_title(s, "5. 개발 진행 현황")
    add_bullets(
        s,
        [
            "완료: 요구사항 확정, 라이브러리 승인, VS 솔루션/프로젝트 생성",
            "완료: config 파서, Fingerprint 검증, SFTP 다운로드 구현",
            "완료: dist 예제, build.bat, 요구사항/서버운용 MD 문서",
            "대기: 개발 PC에서 Release|x64 빌드 검증",
            "대기: 실제 Linux SFTP 연동 테스트 및 배포 패키지 점검",
        ],
    )

    # 7 Structure
    s = new_slide()
    add_title(s, "6. 솔루션 구성")
    add_bullets(
        s,
        [
            "SftpClient.sln",
            "  · libssh2 : WinCNG 정적 라이브러리 (/MT)",
            "  · Client  : Console 다운로드 앱 (/MT)",
            "third_party/libssh2-1.11.1 : vendor 소스",
            "dist/ : Client.exe + config.ini + LICENSE",
            "docs/ : 요구사항 MD, 서버운용 MD, 본 PPT",
        ],
        size=18,
    )

    # 8 Server
    s = new_slide()
    add_title(s, "7. 서버 운용 핵심 요구")
    add_bullets(
        s,
        [
            "OpenSSH SFTP, PasswordAuthentication (PoC)",
            "HostKey: RSA 또는 ECDSA 유지 (Ed25519 only 금지)",
            "Fingerprint를 클라이언트 config.ini에 필수 배포",
            "다운로드 전용 계정 + 최소 권한 (Chroot 권고)",
            "방화벽: 클라이언트 → 서버 TCP 허용, 가능하면 소스 IP 제한",
            "Host Key 재생성 시 전 클라이언트 설정 동시 갱신",
        ],
    )

    # 9 Next
    s = new_slide()
    add_title(s, "8. 다음 단계")
    add_bullets(
        s,
        [
            "Visual Studio에서 Release|x64 빌드 (`scripts\\build.bat`)",
            "서버 fingerprint 산출 후 dist\\config.ini 기입",
            "1~N개 파일 다운로드 E2E 테스트",
            "Runtime 미설치 Windows 10/11 PC에서 실행 확인",
            "운영 전환 시: 키 인증, 설정 평문 최소화, IP 제한 검토",
        ],
    )

    # 10 Close
    s = new_slide()
    box = s.shapes.add_textbox(Inches(0.7), Inches(2.6), Inches(12), Inches(1))
    tf = box.text_frame
    tf.clear()
    p = tf.paragraphs[0]
    run = p.add_run()
    run.text = "상세 문서는 docs 폴더를 참고하세요"
    set_run(run, size=28, bold=True, color=TITLE)

    box = s.shapes.add_textbox(Inches(0.7), Inches(3.7), Inches(12), Inches(1.5))
    tf = box.text_frame
    tf.word_wrap = True
    tf.clear()
    lines = [
        "docs/요구사항_및_개발진행.md",
        "docs/서버_설정_및_운용요구사항.md",
        "BUILD.txt",
    ]
    for i, line in enumerate(lines):
        p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        run = p.add_run()
        run.text = line
        set_run(run, size=18, bold=True, color=ACCENT)
        p.space_after = Pt(6)

    total = len(slides_meta)
    for i, slide in enumerate(slides_meta, start=1):
        add_footer(slide, i, total)

    prs.save(OUTPUT)
    print("Wrote", OUTPUT)


if __name__ == "__main__":
    main()
