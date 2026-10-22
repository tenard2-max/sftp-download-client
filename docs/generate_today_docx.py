# -*- coding: utf-8 -*-
"""2026-10-22 작업 내역과 소스 설명을 워드 문서로 만든다."""

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml.ns import qn
from docx.shared import Cm, Pt, RGBColor

OUT = r"D:\Cursor\SFTP\docs\2026-10-22_작업내역_및_소스설명.docx"


def set_run_font(run, size, bold=False, color=None):
    run.bold = bold
    run.font.size = Pt(size)
    run.font.name = "Malgun Gothic"
    run._element.rPr.rFonts.set(qn("w:eastAsia"), "Malgun Gothic")
    if color is not None:
        run.font.color.rgb = color


def add_heading_text(doc, text, level):
    paragraph = doc.add_heading(text, level=level)
    for run in paragraph.runs:
        set_run_font(run, 16 if level == 1 else 13, bold=True, color=RGBColor(0, 74, 140))
    return paragraph


def add_body(doc, text):
    paragraph = doc.add_paragraph()
    paragraph.paragraph_format.space_after = Pt(6)
    paragraph.paragraph_format.line_spacing = 1.15
    run = paragraph.add_run(text)
    set_run_font(run, 11)
    return paragraph


def add_bullet(doc, text):
    paragraph = doc.add_paragraph(style="List Bullet")
    paragraph.paragraph_format.space_after = Pt(2)
    if paragraph.runs:
        paragraph.runs[0].text = text
        set_run_font(paragraph.runs[0], 11)
    else:
        run = paragraph.add_run(text)
        set_run_font(run, 11)
    return paragraph


def shade_header(cell):
    tc_pr = cell._tc.get_or_add_tcPr()
    shd = tc_pr.makeelement(qn("w:shd"), {
        qn("w:fill"): "E8F2FA",
        qn("w:val"): "clear",
    })
    tc_pr.append(shd)


def add_table(doc, headers, rows):
    table = doc.add_table(rows=1, cols=len(headers))
    table.style = "Table Grid"
    table.autofit = True
    for index, header in enumerate(headers):
        cell = table.rows[0].cells[index]
        cell.text = ""
        paragraph = cell.paragraphs[0]
        run = paragraph.add_run(header)
        set_run_font(run, 10, bold=True, color=RGBColor(0, 74, 140))
        shade_header(cell)
    for row in rows:
        cells = table.add_row().cells
        for index, value in enumerate(row):
            cells[index].text = ""
            run = cells[index].paragraphs[0].add_run(value)
            set_run_font(run, 10)
    doc.add_paragraph()
    return table


def main():
    doc = Document()
    section = doc.sections[0]
    section.top_margin = Cm(2.0)
    section.bottom_margin = Cm(2.0)
    section.left_margin = Cm(2.2)
    section.right_margin = Cm(2.2)

    style = doc.styles["Normal"]
    style.font.name = "Malgun Gothic"
    style.font.size = Pt(11)
    style.element.rPr.rFonts.set(qn("w:eastAsia"), "Malgun Gothic")

    title = doc.add_paragraph()
    title.alignment = WD_ALIGN_PARAGRAPH.CENTER
    title_run = title.add_run("ATEC SFTP 다운로드 클라이언트")
    set_run_font(title_run, 20, bold=True, color=RGBColor(0, 74, 140))

    subtitle = doc.add_paragraph()
    subtitle.alignment = WD_ALIGN_PARAGRAPH.CENTER
    subtitle_run = subtitle.add_run("2026-10-22 작업 내역 및 소스 설명")
    set_run_font(subtitle_run, 14, bold=True)

    date_line = doc.add_paragraph()
    date_line.alignment = WD_ALIGN_PARAGRAPH.CENTER
    date_run = date_line.add_run("작성일 2026-10-22")
    set_run_font(date_run, 11, color=RGBColor(90, 110, 130))

    add_heading_text(doc, "1. 개요", 1)
    add_body(doc, "오늘 작업은 서버 창과 클라이언트 창을 나누고, 접속에 쓰는 설정과 서버 계정 설정을 섞이지 않게 한 것이다. 이어서 로그인 실패 처리, 전송 중 서버 종료, 재시도, 이어 받기, DOS 창 표시, 클라이언트 창의 자동 시작을 반영했다.")
    add_body(doc, "실행 파일은 세 개다. sft_ser.exe 는 서버 기동과 서버 설정, sft_cli.exe 는 수신 진행 화면, Client.exe 는 실제 SFTP 수신이다. 화면 두 개는 같은 Presenter.cpp 를 빌드한 뒤 실행 파일 이름으로 역할을 가른다.")

    add_heading_text(doc, "2. 오늘 작업 내역", 1)

    add_heading_text(doc, "2.1 클라이언트 설정과 서버 설정 분리", 2)
    add_body(doc, "클라이언트 접속은 config.ini 만 읽는다. 서버 설정 창에서 저장한 User, Password, HostFingerprint 는 config_ser.ini 에만 남고, 프로그램이 config.ini 를 덮어쓰지 않는다. 이전에는 클라이언트가 config_ser.ini 값으로 접속 계정을 덮어써서, config.ini 와 로그의 계정이 달랐다.")

    add_heading_text(doc, "2.2 서버 창과 클라이언트 창 분리", 2)
    add_bullet(doc, "sft_ser.exe (제목 ATEC SFTP SERVER): 서버기동/서버종료, 서버 설정, 종료, 로그. SFTP DOWN 과 진행도는 없다. 진행도 자리는 로그가 채운다.")
    add_bullet(doc, "sft_cli.exe (제목 ATEC SFTP DOWN APP): SFTP DOWN, 개별 진행도, 전체 진행도, 종료, 로그. 서버기동, 서버종료, 서버 설정은 없다.")

    add_heading_text(doc, "2.3 로그인 실패", 2)
    add_bullet(doc, "계정 또는 비밀번호가 틀리면 화면에 '로그인 실패 (계정/비밀번호)'가 나온다.")
    add_bullet(doc, "이 실패는 재시도하지 않고 1회에서 바로 종료한다.")
    add_bullet(doc, "지문 불일치도 재시도 없이 종료한다.")

    add_heading_text(doc, "2.4 서버 종료 시 파일 전송 중단", 2)
    add_body(doc, "서버 종료는 새 접속을 받는 sshd 서비스만 멈추면, 이미 붙은 세션이 파일을 계속 보낸다. 그래서 서버 종료 버튼을 누르면 파일을 읽는 sftp-server.exe 와 이미 붙은 sshd 세션을 끝낸 다음 sshd 서비스를 멈춘다. 권한이 없으면 승인 창이 한 번 뜬다. 서버 창에는 '파일 전송 중단'과 '서버 종료 완료'가 남는다.")

    add_heading_text(doc, "2.5 재시도", 2)
    add_bullet(doc, "초기 접속 실패(TCP, SSH 핸드셰이크, SFTP 세션 시작)는 5회까지, 간격 5초. config.ini 의 RetryCount 다. 기본값도 5로 바꿨다.")
    add_bullet(doc, "파일 전송 실패는 20회까지, 그리고 처음 전송을 시도한 시점부터 20분 안에서만 다시 시도한다. 둘 중 하나라도 넘으면 종료한다.")
    add_bullet(doc, "화면에는 접속 단계가 '연결 시도', 수신 단계가 '파일 전송 시도'로 구분된다.")

    add_heading_text(doc, "2.6 이어 받기", 2)
    add_bullet(doc, "로컬 파일 크기가 서버 파일과 같으면 '이미 받음'으로 건너뛴다.")
    add_bullet(doc, "중간에 끊긴 파일은 지우지 않는다. 다음 시도에서 받은 바이트 다음부터 이어 받는다.")
    add_bullet(doc, "로컬 파일이 서버 파일보다 크면 그 파일만 처음부터 다시 받는다.")
    add_bullet(doc, "중간에 실패가 있어도 이어 받기로 모든 파일을 받으면 마지막 표시는 '전체 다운로드 성공'이다. '다운로드 실패'는 끝까지 받지 못하고 종료된 경우에만 나온다.")

    add_heading_text(doc, "2.7 DOS 창", 2)
    add_body(doc, "config.ini 의 DosWindow 로 Client.exe 의 DOS 창을 켜고 끈다. hide 가 기본값이다. show, yes, on, 1 이면 'SFTP Download Client' 창이 열리고 수신 로그가 그 창에도 나온다. Client.log 파일 기록은 어느 쪽이든 남는다. sft_cli.exe 에서 실행할 때도 같은 설정을 따른다.")

    add_heading_text(doc, "2.8 클라이언트 창 자동 시작", 2)
    add_body(doc, "sft_cli.exe 는 창이 열리자마자 SFTP DOWN 을 누른 것처럼 접속과 다운로드를 시작한다. SFTP DOWN 버튼은 보이지만 눌러도 동작하지 않는다. 종료 버튼은 창을 닫는다. 서버 창의 버튼은 그대로 동작한다.")

    add_heading_text(doc, "3. 실행 파일", 1)
    add_table(doc, ["파일", "역할", "만드는 소스"], [
        ["dist\\sft_ser.exe", "서버 기동, 서버 종료, 서버 설정", "Client\\src\\Presenter.cpp"],
        ["dist\\sft_cli.exe", "수신 진행 화면. 시작과 동시에 다운로드", "Client\\src\\Presenter.cpp"],
        ["dist\\Client.exe", "SFTP 접속, 재시도, 이어 받기", "Client\\src\\main.cpp 와 SftpClient.cpp"],
    ])
    add_body(doc, "sft_ser.exe 와 Client.exe 만으로도 서버 기동과 실제 수신을 확인할 수 있다. 진행 막대와 '이어 받기', '전체 다운로드 성공' 문구는 sft_cli.exe 에만 있다. 빌드는 scripts\\build_direct.ps1 이고, 컴파일러는 C:\\Users\\LPH\\mingw64\\bin\\g++.exe 다.")

    add_heading_text(doc, "4. 소스 리스트", 1)
    add_body(doc, "오늘 기능을 담은 소스는 Client\\src 아래 있다. 화면은 Presenter.cpp, 수신 엔진은 SftpClient.cpp 다.")

    add_heading_text(doc, "4.1 화면 — Presenter.cpp", 2)
    add_table(doc, ["함수", "하는 일"], [
        ["WinMain", "창을 만든다. 제목은 서버면 ATEC SFTP SERVER, 클라면 ATEC SFTP DOWN APP"],
        ["IsServerExecutable", "실행 파일 이름이 sft_ser.exe 인지 봐서 서버 창과 클라 창을 가른다"],
        ["CreateControls", "서버는 서버기동과 서버 설정, 클라는 SFTP DOWN 만 만든다. 종료는 둘 다 있다"],
        ["PaintWindow", "로고, 로그 상자, 클라 진행도 막대를 그린다"],
        ["WndProc", "클라 창은 생성 직후 BeginDownload 를 호출한다. SFTP DOWN 클릭은 무시하고 종료는 창을 닫는다"],
        ["BeginDownload / StartDownload", "Client.exe 를 실행하고 Client.log 를 읽어 화면에 반영한다"],
        ["ConsumeImportantLine", "로그를 연결 시도, 파일 전송 시도, 이어 받기, 이미 받음, 성공/실패 문구로 바꾼다"],
        ["StartServerThread", "sshd 서비스를 시작한다. 권한이 없으면 관리자 승인을 띄운다"],
        ["StopServerThread", "sftp-server.exe 와 접속 세션을 끝낸 뒤 sshd 서비스를 멈춘다"],
        ["OpenSetupDialog", "서버 설정 팝업(User, Password, HostFingerprint)을 연다"],
        ["WriteConfigValue", "서버 설정 값을 config_ser.ini 에만 저장한다"],
        ["ApplyAccountUser / ApplyAccountPassword", "Windows 계정에 먼저 반영하고, 성공했을 때만 ini 에 저장한다"],
        ["GenerateHostFingerprint", "새 RSA 호스트 키를 만들어 지문을 config_ser.ini 에 저장한다"],
    ])

    add_heading_text(doc, "4.2 수신 엔진 — SftpClient.cpp", 2)
    add_table(doc, ["함수", "하는 일"], [
        ["RunSftpDownload", "전체 재시도 루프. 초기 접속은 RetryCount 회, 파일 전송은 20회와 20분"],
        ["AttemptDownload", "TCP, 핸드셰이크, 지문, 로그인, SFTP, 파일 수신을 한 번 수행한다"],
        ["AuthenticatePassword", "계정/비밀번호 로그인. 실패하면 AttemptFatal 로 바로 종료"],
        ["DownloadOneFile", "파일 1개 수신. 이미 받은 파일은 건너뛰고, 끊긴 파일은 seek 로 이어 받는다"],
        ["LocalFileSize", "이어 받기 판단에 쓰는 로컬 파일 크기"],
    ])
    add_body(doc, "로그인 실패와 지문 불일치는 AttemptFatal 이다. TCP, 핸드셰이크, SFTP 세션 실패는 초기 접속 재시도고, 파일 읽기 실패는 AttemptTransfer 다.")

    add_heading_text(doc, "4.3 설정과 진입점", 2)
    add_table(doc, ["파일", "하는 일"], [
        ["main.cpp", "Client.exe 진입점. PeekDosWindow 결과로 DOS 창을 연 다음 다운로드를 실행한다"],
        ["Config.h / Config.cpp", "config.ini 를 읽는다. DosWindow, RetryCount 기본 5를 포함한다"],
        ["Logger.cpp", "Client.log 와, DOS 창이 열려 있으면 그 창에도 같은 줄을 남긴다"],
        ["Fingerprint.cpp", "서버 호스트 키의 SHA256 지문을 만들어 config.ini 값과 비교한다"],
        ["Sha256.cpp", "받은 파일의 SHA256 을 계산한다"],
        ["Util.cpp", "경로, 트림, 실행 파일 폴더"],
        ["msvc_compat.h", "MinGW 에서 MSVC 계열 함수 이름을 맞춘다"],
        ["scripts\\build_direct.ps1", "Client.exe, sft_cli.exe, sft_ser.exe 를 직접 빌드한다"],
    ])

    add_heading_text(doc, "4.4 오늘 직접 수정한 파일", 2)
    add_table(doc, ["파일", "수정 내용"], [
        ["Client\\src\\Presenter.cpp", "서버/클라 창 분리, 서버 종료 시 전송 중단, 화면 문구, 클라 자동 시작, 종료 버튼"],
        ["Client\\src\\SftpClient.cpp", "로그인 즉시 종료, 접속 5회/전송 20회·20분, 이어 받기"],
        ["Client\\src\\Config.h / Config.cpp", "RetryCount 기본 5, DosWindow 읽기"],
        ["Client\\src\\main.cpp", "DosWindow=show 일 때 DOS 창 생성"],
        ["scripts\\build_direct.ps1", "sft_cli.exe 와 sft_ser.exe 생성, Client.exe 는 창 없는 서브시스템으로 링크"],
        ["dist\\config.ini", "RetryCount=5, DosWindow=hide. 계정 값은 프로그램이 바꾸지 않는다"],
        ["dist\\config.ini.example", "같은 키의 예시와 설명"],
    ])

    add_heading_text(doc, "5. 설정 파일", 1)
    add_table(doc, ["파일", "누가 읽고 쓰나", "내용"], [
        ["dist\\config.ini", "Client.exe 가 읽기만 한다", "ServerIP, Port, ServerDirectory, User, Password, HostFingerprint, TimeoutSec, RetryCount, RetryIntervalSec, DosWindow, 파일 목록"],
        ["dist\\config_ser.ini", "sft_ser.exe 서버 설정 창이 읽고 쓴다", "User, Password, HostFingerprint. Windows 계정과 호스트 키에 반영한 뒤 저장"],
    ])
    add_body(doc, "클라이언트 로그인에 쓰이는 계정은 config.ini 다. 서버 설정으로 바꾼 계정은 config_ser.ini 와 Windows 계정에 있다. 두 파일의 계정, 비밀번호, 지문이 같아야 접속이 된다.")

    add_heading_text(doc, "6. 확인 방법", 1)
    add_bullet(doc, "서버: dist\\sft_ser.exe 를 실행하고 서버기동, 서버 설정, 서버종료를 본다.")
    add_bullet(doc, "수신 화면: dist\\sft_cli.exe 를 실행하면 바로 다운로드가 시작된다. 종료 버튼으로 창을 닫는다.")
    add_bullet(doc, "수신만: dist 폴더에서 Client.exe 를 실행한다. 결과는 dist\\Client.log 에 남는다.")
    add_bullet(doc, "DOS 창: config.ini 의 DosWindow 를 show 로 바꾼 뒤 Client.exe 또는 sft_cli.exe 를 실행한다.")

    doc.save(OUT)
    print(OUT)


if __name__ == "__main__":
    main()
