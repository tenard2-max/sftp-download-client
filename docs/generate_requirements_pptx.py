# -*- coding: utf-8 -*-
"""Generate requirements & progress PowerPoint at project root."""

from pptx import Presentation
from pptx.util import Inches, Pt
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_SHAPE

OUTPUT = r"D:\Cursor\SFTP\요구사항_및_개발진행.pptx"

BG = RGBColor(0xFF, 0xFF, 0xFF)
ACCENT = RGBColor(0x0B, 0x5F, 0xAD)
TEXT = RGBColor(0x1A, 0x1A, 0x1A)
MUTED = RGBColor(0x44, 0x44, 0x44)
CARD = RGBColor(0xF3, 0xF6, 0xFA)
CARD_LINE = RGBColor(0xC5, 0xD0, 0xDC)
TITLE = RGBColor(0x0A, 0x2E, 0x5C)
WARN = RGBColor(0xB5, 0x5A, 0x00)


def set_run(run, size=18, bold=False, color=TEXT):
    run.font.size = Pt(size)
    run.font.bold = bold
    run.font.color.rgb = color
    run.font.name = "Malgun Gothic"


def add_bg(slide):
    shape = slide.shapes.add_shape(
        MSO_SHAPE.RECTANGLE, Inches(0), Inches(0), Inches(13.333), Inches(7.5)
    )
    shape.fill.solid()
    shape.fill.fore_color.rgb = BG
    shape.line.fill.background()


def add_accent_bar(slide):
    bar = slide.shapes.add_shape(
        MSO_SHAPE.RECTANGLE, Inches(0), Inches(0), Inches(0.18), Inches(7.5)
    )
    bar.fill.solid()
    bar.fill.fore_color.rgb = ACCENT
    bar.line.fill.background()


def add_title(slide, text, top=0.35, size=30):
    box = slide.shapes.add_textbox(Inches(0.55), Inches(top), Inches(12.2), Inches(0.7))
    tf = box.text_frame
    tf.clear()
    p = tf.paragraphs[0]
    run = p.add_run()
    run.text = text
    set_run(run, size=size, bold=True, color=TITLE)


def add_bullets(slide, items, left=0.55, top=1.2, width=12.2, height=5.5, size=16):
    box = slide.shapes.add_textbox(Inches(left), Inches(top), Inches(width), Inches(height))
    tf = box.text_frame
    tf.word_wrap = True
    tf.clear()
    for i, item in enumerate(items):
        p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        run = p.add_run()
        run.text = item
        set_run(run, size=size, color=TEXT)
        p.space_after = Pt(6)


def add_footer(slide, page, total):
    box = slide.shapes.add_textbox(Inches(0.55), Inches(7.05), Inches(12.2), Inches(0.3))
    tf = box.text_frame
    tf.clear()
    p = tf.paragraphs[0]
    run = p.add_run()
    run.text = f"전체 요구사항 및 개발 진행 내역  |  {page}/{total}"
    set_run(run, size=11, color=MUTED)


def add_card(slide, x, y, w, h, title, body_lines, title_color=ACCENT):
    shape = slide.shapes.add_shape(
        MSO_SHAPE.ROUNDED_RECTANGLE, Inches(x), Inches(y), Inches(w), Inches(h)
    )
    shape.fill.solid()
    shape.fill.fore_color.rgb = CARD
    shape.line.color.rgb = CARD_LINE
    shape.line.width = Pt(1.25)
    try:
        shape.adjustments[0] = 0.08
    except Exception:
        pass

    t = slide.shapes.add_textbox(Inches(x + 0.2), Inches(y + 0.15), Inches(w - 0.4), Inches(0.4))
    tf = t.text_frame
    tf.clear()
    p = tf.paragraphs[0]
    run = p.add_run()
    run.text = title
    set_run(run, size=15, bold=True, color=title_color)

    b = slide.shapes.add_textbox(Inches(x + 0.2), Inches(y + 0.55), Inches(w - 0.4), Inches(h - 0.7))
    tf = b.text_frame
    tf.word_wrap = True
    tf.clear()
    for i, line in enumerate(body_lines):
        p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        run = p.add_run()
        run.text = line
        set_run(run, size=13, color=TEXT)
        p.space_after = Pt(3)


def main():
    prs = Presentation()
    prs.slide_width = Inches(13.333)
    prs.slide_height = Inches(7.5)
    blank = prs.slide_layouts[6]
    slides = []

    def new_slide():
        s = prs.slides.add_slide(blank)
        add_bg(s)
        add_accent_bar(s)
        slides.append(s)
        return s

    # 1 Title
    s = new_slide()
    box = s.shapes.add_textbox(Inches(0.7), Inches(2.0), Inches(12), Inches(1.3))
    tf = box.text_frame
    tf.clear()
    p = tf.paragraphs[0]
    run = p.add_run()
    run.text = "전체 요구사항 및 개발 진행 내역"
    set_run(run, size=34, bold=True, color=TITLE)

    box = s.shapes.add_textbox(Inches(0.7), Inches(3.4), Inches(12), Inches(1.0))
    tf = box.text_frame
    tf.clear()
    p = tf.paragraphs[0]
    run = p.add_run()
    run.text = "Windows SFTP Download Client PoC"
    set_run(run, size=22, color=MUTED)

    box = s.shapes.add_textbox(Inches(0.7), Inches(5.0), Inches(12), Inches(0.8))
    tf = box.text_frame
    tf.clear()
    p = tf.paragraphs[0]
    run = p.add_run()
    run.text = "2026-07-29  |  v1.0  |  소스/솔루션 구성 완료 · VS 빌드 검증 대기"
    set_run(run, size=15, bold=True, color=ACCENT)

    # 2 Purpose
    s = new_slide()
    add_title(s, "1. 프로젝트 목적")
    add_bullets(
        s,
        [
            "Linux SFTP 서버 → Windows 로컬로 사전 지정 파일 자동 다운로드",
            "파일 개수: config.ini 기준 1~20개",
            "사용자 조작 없음 (GUI / 목록 조회 / 업로드 / 선택 없음)",
            "배포: Client.exe + config.ini (+ LICENSE), 목표 DLL 없음 · 10MB 미만",
            "고객 PC Runtime·별도 SW 설치 금지 → /MT 정적 CRT",
            "Visual Studio 미설치 PC에서도 실행 가능해야 함",
        ],
    )

    # 3 Dev env + Deploy
    s = new_slide()
    add_title(s, "2. 개발 환경 · 배포 조건")
    add_card(s, 0.55, 1.2, 6.0, 5.3, "개발 환경", [
        "언어: C++ Console",
        "IDE: VS2015 (v140) ※ VS2008→상향 확정",
        "플랫폼: Windows 10/11 x64",
        "제품에 C#/.NET/Python/Java/Node 금지",
        "CRT: /MT (Runtime 미설치 지원)",
    ])
    add_card(s, 6.8, 1.2, 5.9, 5.3, "배포 조건", [
        "정적 링크 우선",
        "불가 시 필요 DLL만 동폴더 배포",
        "목표: Client.exe",
        "         config.ini",
        "         LICENSE_libssh2.txt",
        "설계: libssh2+WinCNG → DLL 없음",
    ])

    # 4 Flow
    s = new_slide()
    add_title(s, "3. 기능 요구사항 — 12단계 자동 흐름")
    add_bullets(
        s,
        [
            "1) config.ini 읽기  →  2) TCP/SSH 연결  →  3) Host Key 수신",
            "4) SHA-256 Fingerprint 계산  →  5) 설정값 비교 (불일치/미등록 시 종료)",
            "6) ID/PW 인증  →  7) SFTP 세션 시작",
            "8) RemoteFile 존재 확인  →  9) LocalFile 다운로드",
            "10) 성공 검증  →  11) 연결 정상 종료  →  12) 프로그램 종료",
            "파일 키: RemoteFile/LocalFile 또는 RemoteFileN/LocalFileN (1~20)",
        ],
        size=16,
    )

    # 5 Security + Errors
    s = new_slide()
    add_title(s, "4. 보안 · 오류 처리")
    add_card(s, 0.55, 1.2, 6.0, 5.3, "보안", [
        "SSH/SFTP 직접 구현 금지",
        "검증된 라이브러리 사용",
        "Host Key 검증 생략 금지",
        "trust-all 코드 작성 금지",
        "Fingerprint 불일치 → 즉시 종료",
        "미등록 → 연결 금지",
    ], title_color=WARN)
    add_card(s, 6.8, 1.2, 5.9, 5.3, "콘솔 오류 메시지", [
        "서버 연결 실패",
        "Host Key 불일치 / 미등록",
        "로그인 실패",
        "RemoteFile 없음",
        "LocalFile 생성 실패",
        "다운로드 실패 / 성공",
    ])

    # 6 Library
    s = new_slide()
    add_title(s, "5. 라이브러리 선정")
    add_bullets(
        s,
        [
            "선정: libssh2 1.11.1 + WinCNG (정적 링크) — BSD-3-Clause",
            "이유: 정적 링크, OpenSSL DLL 불필요, SHA256 Host Key API, VS2015",
            "제외: libssh(LGPL), libcurl(의존성↑), OpenSSH 바이너리 호출",
            "제한: Ed25519 host key 미지원 → 서버 RSA/ECDSA 필요",
            "PoC: LIBSSH2_ECDSA_WINCNG 활성화 (Windows 10/11)",
        ],
    )

    # 7 Progress timeline
    s = new_slide()
    add_title(s, "6. 개발 진행 타임라인")
    add_bullets(
        s,
        [
            "1. 요구 분석·라이브러리 선정 → libssh2+WinCNG 제안",
            "2. IDE 상향 협의 → VS2015 확정",
            "3. 파일 개수 협의 → 1~20개 확정",
            "4. 라이브러리/빌드 승인 → 구현 착수",
            "5~7. libssh2 vendor · 솔루션 생성 · Client 소스 구현",
            "8~9. dist/build 스크립트 · MD/PPT 문서화",
        ],
        size=16,
    )

    # 8 Modules + Build
    s = new_slide()
    add_title(s, "7. 구현 모듈 · 빌드 설정")
    add_card(s, 0.55, 1.2, 6.0, 5.3, "구현 모듈", [
        "Config — ini 파싱/검증",
        "Fingerprint — SHA256 비교",
        "SftpClient — 연결·인증·다운로드",
        "Util — 경로/디렉터리 유틸",
        "libssh2 프로젝트 — WinCNG 정적 lib",
    ])
    add_card(s, 6.8, 1.2, 5.9, 5.3, "빌드", [
        "Release|x64 / Toolset v140",
        "Runtime: /MT (/MTd Debug)",
        "링크: ws2_32, bcrypt, crypt32,",
        "      ntdll, shlwapi",
        "출력: build\\Release\\Client.exe",
        "      → dist\\Client.exe",
    ])

    # 9 Structure
    s = new_slide()
    add_title(s, "8. 솔루션 / 디렉터리 구조")
    add_bullets(
        s,
        [
            "SftpClient.sln",
            "  · Client\\ — Console 앱 소스",
            "  · libssh2\\ — 정적 라이브러리 프로젝트",
            "third_party\\libssh2-1.11.1 — vendor 소스",
            "dist\\ — Client.exe, config.ini, LICENSE",
            "docs\\ — 요구사항 MD, 서버운용 MD",
            "scripts\\build.bat, BUILD.txt",
        ],
        size=16,
    )

    # 10 Remaining + Next
    s = new_slide()
    add_title(s, "9. 미완료 · 다음 작업")
    add_bullets(
        s,
        [
            "대기: 개발 PC에서 Release|x64 MSBuild 빌드 검증",
            "대기: 실제 Linux SFTP 연동 E2E 테스트",
            "대기: Runtime 미설치 Win10/11에서 실행 확인",
            "대기: 배포 패키지 용량(10MB 미만) 최종 점검",
            "빌드: scripts\\build.bat 또는 VS에서 Release|x64",
            "Fingerprint: ssh-keygen -lf ... -E sha256 → config.ini",
        ],
    )

    # 11 Close
    s = new_slide()
    box = s.shapes.add_textbox(Inches(0.7), Inches(2.4), Inches(12), Inches(0.8))
    tf = box.text_frame
    tf.clear()
    p = tf.paragraphs[0]
    run = p.add_run()
    run.text = "원본 / 관련 문서"
    set_run(run, size=28, bold=True, color=TITLE)

    box = s.shapes.add_textbox(Inches(0.7), Inches(3.4), Inches(12), Inches(2.2))
    tf = box.text_frame
    tf.word_wrap = True
    tf.clear()
    for i, line in enumerate([
        "docs/요구사항_및_개발진행.md",
        "docs/서버_설정_및_운용요구사항.md",
        "SFTP_PoC_개요.pptx",
        "서버_설정_및_운용요구사항.pptx",
        "BUILD.txt",
    ]):
        p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        run = p.add_run()
        run.text = line
        set_run(run, size=17, bold=True, color=ACCENT)
        p.space_after = Pt(5)

    total = len(slides)
    for i, slide in enumerate(slides, start=1):
        add_footer(slide, i, total)

    prs.save(OUTPUT)
    print("Wrote", OUTPUT)


if __name__ == "__main__":
    main()
