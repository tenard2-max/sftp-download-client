# -*- coding: utf-8 -*-
"""Generate server configuration & operations PowerPoint at project root."""

from pptx import Presentation
from pptx.util import Inches, Pt
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_SHAPE

OUTPUT = r"D:\Cursor\SFTP\서버_설정_및_운용요구사항.pptx"

# Light (white) background + high-contrast text for readability
BG = RGBColor(0xFF, 0xFF, 0xFF)
ACCENT = RGBColor(0x0B, 0x5F, 0xAD)
TEXT = RGBColor(0x1A, 0x1A, 0x1A)
MUTED = RGBColor(0x44, 0x44, 0x44)
CARD = RGBColor(0xF3, 0xF6, 0xFA)
CARD_LINE = RGBColor(0xC5, 0xD0, 0xDC)
TITLE = RGBColor(0x0A, 0x2E, 0x5C)
WARN = RGBColor(0xB5, 0x5A, 0x00)


def set_run(run, size=18, bold=False, color=TEXT):
    run.font.size = Pt(size)
    run.font.bold = bold
    run.font.color.rgb = color
    run.font.name = "Malgun Gothic"


def add_bg(slide):
    shape = slide.shapes.add_shape(
        MSO_SHAPE.RECTANGLE, Inches(0), Inches(0), Inches(13.333), Inches(7.5)
    )
    shape.fill.solid()
    shape.fill.fore_color.rgb = BG
    shape.line.fill.background()


def add_accent_bar(slide):
    bar = slide.shapes.add_shape(
        MSO_SHAPE.RECTANGLE, Inches(0), Inches(0), Inches(0.18), Inches(7.5)
    )
    bar.fill.solid()
    bar.fill.fore_color.rgb = ACCENT
    bar.line.fill.background()


def add_title(slide, text, top=0.35, size=30):
    box = slide.shapes.add_textbox(Inches(0.55), Inches(top), Inches(12.2), Inches(0.7))
    tf = box.text_frame
    tf.clear()
    p = tf.paragraphs[0]
    run = p.add_run()
    run.text = text
    set_run(run, size=size, bold=True, color=TITLE)


def add_bullets(slide, items, left=0.55, top=1.2, width=12.2, height=5.5, size=17):
    box = slide.shapes.add_textbox(Inches(left), Inches(top), Inches(width), Inches(height))
    tf = box.text_frame
    tf.word_wrap = True
    tf.clear()
    for i, item in enumerate(items):
        p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        run = p.add_run()
        run.text = item
        set_run(run, size=size, color=TEXT)
        p.space_after = Pt(7)


def add_footer(slide, page, total):
    box = slide.shapes.add_textbox(Inches(0.55), Inches(7.05), Inches(12.2), Inches(0.3))
    tf = box.text_frame
    tf.clear()
    p = tf.paragraphs[0]
    run = p.add_run()
    run.text = f"SFTP 서버 설정 및 운용 요구사항  |  {page}/{total}"
    set_run(run, size=11, color=MUTED)


def add_card(slide, x, y, w, h, title, body_lines, title_color=ACCENT):
    shape = slide.shapes.add_shape(
        MSO_SHAPE.ROUNDED_RECTANGLE, Inches(x), Inches(y), Inches(w), Inches(h)
    )
    shape.fill.solid()
    shape.fill.fore_color.rgb = CARD
    shape.line.color.rgb = CARD_LINE
    shape.line.width = Pt(1.25)
    try:
        shape.adjustments[0] = 0.08
    except Exception:
        pass

    t = slide.shapes.add_textbox(Inches(x + 0.2), Inches(y + 0.15), Inches(w - 0.4), Inches(0.4))
    tf = t.text_frame
    tf.clear()
    p = tf.paragraphs[0]
    run = p.add_run()
    run.text = title
    set_run(run, size=16, bold=True, color=title_color)

    b = slide.shapes.add_textbox(Inches(x + 0.2), Inches(y + 0.55), Inches(w - 0.4), Inches(h - 0.7))
    tf = b.text_frame
    tf.word_wrap = True
    tf.clear()
    for i, line in enumerate(body_lines):
        p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        run = p.add_run()
        run.text = line
        set_run(run, size=13, color=TEXT)
        p.space_after = Pt(3)


def main():
    prs = Presentation()
    prs.slide_width = Inches(13.333)
    prs.slide_height = Inches(7.5)
    blank = prs.slide_layouts[6]
    slides = []

    def new_slide():
        s = prs.slides.add_slide(blank)
        add_bg(s)
        add_accent_bar(s)
        slides.append(s)
        return s

    # 1 Title
    s = new_slide()
    box = s.shapes.add_textbox(Inches(0.7), Inches(2.1), Inches(12), Inches(1.2))
    tf = box.text_frame
    tf.clear()
    p = tf.paragraphs[0]
    run = p.add_run()
    run.text = "SFTP 서버 기본 설정 및 운용 요구사항"
    set_run(run, size=36, bold=True, color=TITLE)

    box = s.shapes.add_textbox(Inches(0.7), Inches(3.5), Inches(12), Inches(1.0))
    tf = box.text_frame
    tf.clear()
    p = tf.paragraphs[0]
    run = p.add_run()
    run.text = "Linux OpenSSH 서버 기준  ·  Windows Client.exe PoC 연동"
    set_run(run, size=20, color=MUTED)

    box = s.shapes.add_textbox(Inches(0.7), Inches(5.2), Inches(12), Inches(0.6))
    tf = box.text_frame
    tf.clear()
    p = tf.paragraphs[0]
    run = p.add_run()
    run.text = "2026-07-29  |  문서 v1.0  |  docs/서버_설정_및_운용요구사항.md"
    set_run(run, size=15, bold=True, color=ACCENT)

    # 2 Purpose
    s = new_slide()
    add_title(s, "1. 목적")
    add_bullets(
        s,
        [
            "Windows 클라이언트가 약속된 파일만 안정적으로 다운로드하도록 서버 기준을 정의",
            "계정 · 권한 · Host Key · 네트워크 · 모니터링 기준 포함",
            "PoC 단계에서도 운영과 동일한 보안 원칙 적용",
            "  · Host Key(Fingerprint) 검증",
            "  · 최소 권한 계정",
            "상세 MD: docs/서버_설정_및_운용요구사항.md",
        ],
    )

    # 3 Prerequisites
    s = new_slide()
    add_title(s, "2. 서버 전제 조건")
    add_card(s, 0.55, 1.25, 4.0, 5.2, "플랫폼", [
        "OS: Linux (2021년경 신형)",
        "SSH/SFTP: OpenSSH sshd",
        "포트: 기본 22",
        "변경 시 Client Port와 동일",
    ])
    add_card(s, 4.75, 1.25, 4.0, 5.2, "인증 / Host Key", [
        "PoC: ID/PW 비밀번호 인증",
        "Host Key: RSA 또는 ECDSA 필수",
        "Ed25519 only 구성 금지",
        "(Client WinCNG 미지원)",
    ], title_color=WARN)
    add_card(s, 8.95, 1.25, 3.8, 5.2, "네트워크", [
        "클라이언트 → 서버 TCP 허용",
        "가능하면 소스 IP 제한",
        "Host/DNS 해석 가능해야 함",
        "방화벽 idle timeout 고려",
    ])

    # 4 sshd
    s = new_slide()
    add_title(s, "3. 기본 sshd 설정 권고")
    add_bullets(
        s,
        [
            "설정: /etc/ssh/sshd_config (또는 sshd_config.d/*.conf)",
            "HostKey rsa / ecdsa 활성화 권장 (ed25519는 있어도 무방)",
            "PasswordAuthentication yes  (PoC)  /  PermitEmptyPasswords no",
            "PermitRootLogin no  /  Subsystem sftp internal-sftp",
            "변경 후: sudo sshd -t  →  systemctl reload sshd",
            "운영 전환 시 공개키 인증으로 전환 검토",
        ],
    )

    # 5 Chroot
    s = new_slide()
    add_title(s, "4. SFTP 전용 계정 + Chroot (권고)")
    add_bullets(
        s,
        [
            "다운로드 전용 계정은 SFTP only + Chroot 권장",
            "Match User sftpdownload",
            "  · ChrootDirectory /data/sftp/%u",
            "  · ForceCommand internal-sftp",
            "  · AllowTcpForwarding no / X11Forwarding no",
            "디렉터리: chroot root=755(root), data 디렉터리=계정 소유",
            "RemoteFile는 chroot 기준 경로 (예: /data/test.dat)",
            "Chroot 미사용 PoC면 서버 절대경로 그대로 사용",
        ],
        size=16,
    )

    # 6 Fingerprint
    s = new_slide()
    add_title(s, "5. Host Fingerprint 운용")
    add_bullets(
        s,
        [
            "클라이언트는 config.ini HostFingerprint와 SHA256을 반드시 비교",
            "미등록 · 불일치 시 연결 금지 (trust-all 금지)",
            "확인: ssh-keygen -lf /etc/ssh/ssh_host_ecdsa_key.pub -E sha256",
            "config.ini 형식: HostFingerprint=SHA256:....",
            "Host Key 재생성 시 모든 클라이언트 설정 동시 배포·사전 공지",
            "Ed25519만 남기는 구성 금지 — RSA/ECDSA 유지",
        ],
    )

    # 7 Account/Files
    s = new_slide()
    add_title(s, "6. 계정 · 권한 · 파일 배치")
    add_card(s, 0.55, 1.25, 6.0, 5.2, "계정", [
        "다운로드 전용 (업로드 불필요 시 비활성)",
        "sudo 권한 부여 금지",
        "복잡 비밀번호 + 정기 변경",
        "약속된 디렉터리만 접근",
        "가능하면 소스 IP 제한",
    ])
    add_card(s, 6.8, 1.25, 5.9, 5.2, "파일", [
        "사전 합의 1~20개",
        "RemoteFile 경로와 완전 일치",
        "일반 파일만 (디렉터리 불가)",
        "계정에 읽기 권한",
        "원자적 배치: *.part → rename",
        "디스크/inode 임계 모니터링",
    ])

    # 8 Network / Log
    s = new_slide()
    add_title(s, "7. 네트워크 · 로깅 · 변경관리")
    add_bullets(
        s,
        [
            "방화벽: 클라이언트 → 서버 TCP Port 허용, 소스 IP allow 권고",
            "로그: /var/log/auth.log 또는 /var/log/secure 정기 점검",
            "모니터링: 인증 실패 급증, 디스크, sshd down, Host Key 변경",
            "Host Key·sshd_config·비밀번호 변경은 변경관리 + 클라이언트 동시 배포",
            "클라이언트 오류 ↔ 서버 auth 로그 대조로 원인 분리",
        ],
    )

    # 9 Checklist
    s = new_slide()
    add_title(s, "8. 배포 전 체크리스트")
    add_bullets(
        s,
        [
            "□ sshd 기동 / 포트 listen",
            "□ RSA 또는 ECDSA host key 존재",
            "□ SHA256 fingerprint 산출 및 클라이언트 배포",
            "□ 다운로드 계정으로 SFTP 수동 접속 성공",
            "□ RemoteFile 경로·권한·존재 확인",
            "□ 방화벽에서 클라이언트 IP 허용",
            "□ Chroot 사용 시 경로 매핑 문서화",
            "□ NTP 시각 동기화 / 로그 로테이션",
        ],
        size=16,
    )

    # 10 Ops beyond PoC + Incident
    s = new_slide()
    add_title(s, "9. 운영 고도화 · 장애 대응")
    add_card(s, 0.55, 1.25, 6.0, 5.2, "PoC 이후 권고", [
        "공개키 인증 전환",
        "config 평문 최소화 / ACL",
        "VPN 또는 IP 제한",
        "읽기 전용 + 디렉터리 분리",
        "감사 로그 SIEM 연동",
    ])
    add_card(s, 6.8, 1.25, 5.9, 5.2, "장애 1차 조치", [
        "전체 실패 → sshd/디스크/방화벽",
        "특정 PC → IP allow / config.ini",
        "Key 불일치 → Host Key 변경 여부",
        "로그인 실패 → 계정 lock/PAM",
        "RemoteFile 없음 → 경로/배치",
    ], title_color=WARN)

    # 11 Close
    s = new_slide()
    box = s.shapes.add_textbox(Inches(0.7), Inches(2.5), Inches(12), Inches(1))
    tf = box.text_frame
    tf.clear()
    p = tf.paragraphs[0]
    run = p.add_run()
    run.text = "원본 상세 문서"
    set_run(run, size=28, bold=True, color=TITLE)

    box = s.shapes.add_textbox(Inches(0.7), Inches(3.6), Inches(12), Inches(1.8))
    tf = box.text_frame
    tf.word_wrap = True
    tf.clear()
    for i, line in enumerate([
        "docs/서버_설정_및_운용요구사항.md",
        "docs/요구사항_및_개발진행.md",
        "SFTP_PoC_개요.pptx",
    ]):
        p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        run = p.add_run()
        run.text = line
        set_run(run, size=18, bold=True, color=ACCENT)
        p.space_after = Pt(6)

    total = len(slides)
    for i, slide in enumerate(slides, start=1):
        add_footer(slide, i, total)

    prs.save(OUTPUT)
    print("Wrote", OUTPUT)


if __name__ == "__main__":
    main()
